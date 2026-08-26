"""Path-confined Word and PDF tools with structural post-write validation."""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any, Literal

from pydantic import BaseModel, ConfigDict, Field

from agent.tools.base import ToolSpec
from agent.tools.registry import ToolRegistry

_SAFE_FILENAME = re.compile(r"^[\w\-. ]{1,128}$", re.UNICODE)


class DocumentSection(BaseModel):
    """One section accepted by document creation tools."""

    model_config = ConfigDict(extra="forbid")

    heading: str = Field(min_length=1, max_length=160)
    body: str = Field(default="", max_length=8000)
    bullets: list[str] = Field(default_factory=list, max_length=20)


class MetricRow(BaseModel):
    """One compact label/value row in a report."""

    model_config = ConfigDict(extra="forbid")

    label: str = Field(min_length=1, max_length=120)
    value: str = Field(min_length=1, max_length=240)


class CreateDocumentArguments(BaseModel):
    """Shared structured input for DOCX and PDF creation."""

    model_config = ConfigDict(extra="forbid")

    filename: str = Field(min_length=1, max_length=128)
    title: str = Field(min_length=1, max_length=200)
    subtitle: str = Field(default="", max_length=300)
    author: str = Field(default="李林知", max_length=100)
    sections: list[DocumentSection] = Field(min_length=1, max_length=24)
    metrics: list[MetricRow] = Field(default_factory=list, max_length=30)
    citations: list[str] = Field(default_factory=list, max_length=40)
    overwrite: bool = False


class ReadDocumentArguments(BaseModel):
    """Input for confined document extraction."""

    model_config = ConfigDict(extra="forbid")

    filename: str = Field(min_length=1, max_length=128)
    max_characters: int = Field(default=20_000, ge=100, le=100_000)


class SummarizeDocumentArguments(ReadDocumentArguments):
    """Input for deterministic extractive summaries."""

    summary_characters: int = Field(default=1200, ge=100, le=5000)


class ExportReportArguments(BaseModel):
    """Input for a compact formal metrics report."""

    model_config = ConfigDict(extra="forbid")

    filename: str = Field(min_length=1, max_length=128)
    format: Literal["docx", "pdf"]
    title: str = Field(min_length=1, max_length=200)
    author: str = Field(default="李林知", max_length=100)
    summary: str = Field(min_length=1, max_length=4000)
    metrics: list[MetricRow] = Field(min_length=1, max_length=30)
    findings: list[str] = Field(min_length=1, max_length=30)
    citations: list[str] = Field(default_factory=list, max_length=40)
    overwrite: bool = False


@dataclass(frozen=True)
class DocumentWorkspace:
    """Confine all document reads and writes to one explicit directory."""

    root: Path
    max_file_bytes: int = 10_000_000

    def __post_init__(self) -> None:
        resolved = self.root.resolve()
        resolved.mkdir(parents=True, exist_ok=True)
        object.__setattr__(self, "root", resolved)
        if self.max_file_bytes <= 0:
            raise ValueError("max_file_bytes must be positive.")

    def create_docx(self, **payload: Any) -> dict[str, Any]:
        """Create and structurally validate a styled DOCX report."""
        arguments = CreateDocumentArguments.model_validate(payload)
        path = self._output_path(
            arguments.filename,
            suffix=".docx",
            overwrite=arguments.overwrite,
        )
        self._write_docx(path, arguments)
        extracted = self.read_docx(path.name, max_characters=100_000)
        if arguments.title not in extracted["text"]:
            path.unlink(missing_ok=True)
            raise ValueError("DOCX post-write validation did not find the title.")
        return self._created_payload(path, extracted)

    def create_pdf(self, **payload: Any) -> dict[str, Any]:
        """Create and structurally validate a styled PDF report."""
        arguments = CreateDocumentArguments.model_validate(payload)
        path = self._output_path(
            arguments.filename,
            suffix=".pdf",
            overwrite=arguments.overwrite,
        )
        self._write_pdf(path, arguments)
        extracted = self.read_pdf(path.name, max_characters=100_000)
        if arguments.title not in extracted["text"]:
            path.unlink(missing_ok=True)
            raise ValueError("PDF post-write validation did not find the title.")
        return self._created_payload(path, extracted)

    def read_docx(self, filename: str, max_characters: int = 20_000) -> dict[str, Any]:
        """Extract paragraphs and table cells from a workspace DOCX."""
        from docx import Document

        path = self._input_path(filename, suffix=".docx")
        document = Document(str(path))
        blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
        for table in document.tables:
            for row in table.rows:
                blocks.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(blocks)
        return {
            "filename": path.name,
            "text": text[:max_characters],
            "characters": len(text),
            "paragraphs": len(document.paragraphs),
            "tables": len(document.tables),
            "truncated": len(text) > max_characters,
        }

    def read_pdf(self, filename: str, max_characters: int = 20_000) -> dict[str, Any]:
        """Extract text from every page of a workspace PDF."""
        from pypdf import PdfReader

        path = self._input_path(filename, suffix=".pdf")
        reader = PdfReader(path)
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        return {
            "filename": path.name,
            "text": text[:max_characters],
            "characters": len(text),
            "pages": len(reader.pages),
            "truncated": len(text) > max_characters,
        }

    def summarize_document(
        self,
        filename: str,
        max_characters: int = 20_000,
        summary_characters: int = 1200,
    ) -> dict[str, Any]:
        """Create a deterministic extractive summary without inventing content."""
        suffix = Path(filename).suffix.lower()
        if suffix == ".docx":
            extracted = self.read_docx(filename, max_characters=max_characters)
        elif suffix == ".pdf":
            extracted = self.read_pdf(filename, max_characters=max_characters)
        else:
            raise ValueError("summarize_document supports only .docx and .pdf.")
        paragraphs = [
            line.strip() for line in extracted["text"].splitlines() if line.strip()
        ]
        selected: list[str] = []
        size = 0
        for paragraph in paragraphs:
            if size + len(paragraph) > summary_characters and selected:
                break
            selected.append(paragraph)
            size += len(paragraph)
        return {
            "filename": filename,
            "summary": "\n".join(selected)[:summary_characters],
            "source_characters": extracted["characters"],
            "extractive": True,
        }

    def export_report(self, **payload: Any) -> dict[str, Any]:
        """Export a formal report through the same validated creation path."""
        arguments = ExportReportArguments.model_validate(payload)
        sections = [
            DocumentSection(heading="执行摘要", body=arguments.summary),
            DocumentSection(heading="主要发现", bullets=arguments.findings),
            DocumentSection(
                heading="方法与边界",
                body=(
                    "本报告由受控文档工具生成；所有输出限定在配置的工作目录内，"
                    "并在写入后重新读取关键字段。"
                ),
            ),
        ]
        create_payload = CreateDocumentArguments(
            filename=arguments.filename,
            title=arguments.title,
            subtitle=f"项目验收报告 | {date.today().isoformat()}",
            author=arguments.author,
            sections=sections,
            metrics=arguments.metrics,
            citations=arguments.citations,
            overwrite=arguments.overwrite,
        ).model_dump()
        if arguments.format == "docx":
            return self.create_docx(**create_payload)
        return self.create_pdf(**create_payload)

    def _input_path(self, filename: str, *, suffix: str) -> Path:
        path = self._safe_path(filename, suffix=suffix)
        if not path.is_file():
            raise ValueError(f"Document does not exist: {path.name}")
        if path.stat().st_size > self.max_file_bytes:
            raise ValueError("Document exceeds the configured size limit.")
        return path

    def _output_path(self, filename: str, *, suffix: str, overwrite: bool) -> Path:
        path = self._safe_path(filename, suffix=suffix)
        if path.exists() and not overwrite:
            raise ValueError("Document already exists; set overwrite=true explicitly.")
        return path

    def _safe_path(self, filename: str, *, suffix: str) -> Path:
        if not _SAFE_FILENAME.fullmatch(filename):
            raise ValueError("filename contains unsupported characters.")
        candidate = Path(filename)
        if (
            candidate.is_absolute()
            or candidate.name != filename
            or candidate.suffix.lower() != suffix
        ):
            raise ValueError(
                f"filename must be a plain {suffix} name without directories."
            )
        resolved = (self.root / candidate).resolve()
        if resolved.parent != self.root:
            raise ValueError("Document path escapes the configured workspace.")
        return resolved

    @staticmethod
    def _created_payload(path: Path, extracted: dict[str, Any]) -> dict[str, Any]:
        payload = {
            "filename": path.name,
            "bytes": path.stat().st_size,
            "validated": True,
        }
        for key in ("pages", "paragraphs", "tables"):
            if key in extracted:
                payload[key] = extracted[key]
        return payload

    def _write_docx(self, path: Path, arguments: CreateDocumentArguments) -> None:
        from docx import Document
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor

        document = Document()
        section = document.sections[0]
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = section.bottom_margin = Inches(1)
        section.left_margin = section.right_margin = Inches(1)
        section.header_distance = section.footer_distance = Inches(0.492)

        styles = document.styles
        normal = styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(11)
        normal.paragraph_format.space_after = Pt(8)
        normal.paragraph_format.line_spacing = 1.167
        for name, size, before, after in (
            ("Heading 1", 16, 16, 8),
            ("Heading 2", 13, 12, 6),
            ("Heading 3", 12, 8, 4),
        ):
            style = styles[name]
            style.font.name = "Calibri"
            style.font.size = Pt(size)
            style.font.color.rgb = RGBColor(46, 116, 181)
            style.paragraph_format.space_before = Pt(before)
            style.paragraph_format.space_after = Pt(after)

        header = section.header.paragraphs[0]
        header.text = "MINIGPT AGENT | CONTROLLED DOCUMENT EXPORT"
        header.runs[0].font.size = Pt(8)
        header.runs[0].font.color.rgb = RGBColor(100, 116, 139)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer.add_run(f"{arguments.author}  |  ")
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        footer._p.append(field)

        title = document.add_paragraph()
        title.paragraph_format.space_before = Pt(18)
        title.paragraph_format.space_after = Pt(5)
        run = title.add_run(arguments.title)
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(24)
        run.font.color.rgb = RGBColor(15, 43, 76)
        if arguments.subtitle:
            subtitle = document.add_paragraph(arguments.subtitle)
            subtitle.paragraph_format.space_after = Pt(14)
            subtitle.runs[0].font.size = Pt(12)
            subtitle.runs[0].font.color.rgb = RGBColor(71, 85, 105)
        meta = document.add_paragraph()
        meta.add_run("作者：").bold = True
        meta.add_run(arguments.author)
        meta.add_run("    日期：").bold = True
        meta.add_run(date.today().isoformat())

        document.add_heading("目录", level=1)
        for item in arguments.sections:
            paragraph = document.add_paragraph(style="List Number")
            paragraph.paragraph_format.left_indent = Inches(0.5)
            paragraph.paragraph_format.first_line_indent = Inches(-0.25)
            paragraph.paragraph_format.space_after = Pt(8)
            paragraph.paragraph_format.line_spacing = 1.167
            paragraph.add_run(item.heading)
        if arguments.metrics:
            paragraph = document.add_paragraph(style="List Number")
            paragraph.add_run("关键指标")
        if arguments.citations:
            paragraph = document.add_paragraph(style="List Number")
            paragraph.add_run("来源与引用")

        for item in arguments.sections:
            document.add_heading(item.heading, level=1)
            if item.body:
                document.add_paragraph(item.body)
            for bullet in item.bullets:
                paragraph = document.add_paragraph(style="List Bullet")
                paragraph.paragraph_format.left_indent = Inches(0.5)
                paragraph.paragraph_format.first_line_indent = Inches(-0.25)
                paragraph.paragraph_format.space_after = Pt(8)
                paragraph.paragraph_format.line_spacing = 1.167
                paragraph.add_run(bullet)

        if arguments.metrics:
            document.add_page_break()
            document.add_heading("关键指标", level=1)
            table = document.add_table(rows=1, cols=2)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.LEFT
            table.autofit = False
            headers = table.rows[0].cells
            headers[0].text, headers[1].text = "指标", "实测值"
            repeat_header = OxmlElement("w:tblHeader")
            repeat_header.set(qn("w:val"), "true")
            table.rows[0]._tr.get_or_add_trPr().append(repeat_header)
            for metric in arguments.metrics:
                cells = table.add_row().cells
                cells[0].text, cells[1].text = metric.label, metric.value
            widths = (3600, 5760)
            self._set_table_geometry(table, widths)
            for row_index, row in enumerate(table.rows):
                for cell in row.cells:
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_before = Pt(0)
                        paragraph.paragraph_format.space_after = Pt(0)
                        paragraph.paragraph_format.line_spacing = 1.0
                    if row_index == 0:
                        cell.paragraphs[0].runs[0].bold = True
                        shading = OxmlElement("w:shd")
                        shading.set(qn("w:fill"), "F2F4F7")
                        cell._tc.get_or_add_tcPr().append(shading)

        if arguments.citations:
            document.add_heading("来源与引用", level=1)
            for citation in arguments.citations:
                paragraph = document.add_paragraph(style="List Bullet")
                paragraph.add_run(citation)

        for paragraph in document.paragraphs:
            for run in paragraph.runs:
                run._element.get_or_add_rPr().get_or_add_rFonts().set(
                    qn("w:eastAsia"), "Microsoft YaHei"
                )
        document.save(str(path))

    @staticmethod
    def _set_table_geometry(table: Any, widths: tuple[int, ...]) -> None:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        table_properties = table._tbl.tblPr
        width = table_properties.first_child_found_in("w:tblW")
        if width is None:
            width = OxmlElement("w:tblW")
            table_properties.append(width)
        width.set(qn("w:w"), str(sum(widths)))
        width.set(qn("w:type"), "dxa")
        indent = OxmlElement("w:tblInd")
        indent.set(qn("w:w"), "120")
        indent.set(qn("w:type"), "dxa")
        table_properties.append(indent)
        cell_margins = OxmlElement("w:tblCellMar")
        for edge, value in (
            ("top", 80),
            ("bottom", 80),
            ("start", 120),
            ("end", 120),
        ):
            margin = OxmlElement(f"w:{edge}")
            margin.set(qn("w:w"), str(value))
            margin.set(qn("w:type"), "dxa")
            cell_margins.append(margin)
        table_properties.append(cell_margins)
        grid = table._tbl.tblGrid
        for child in list(grid):
            grid.remove(child)
        for value in widths:
            column = OxmlElement("w:gridCol")
            column.set(qn("w:w"), str(value))
            grid.append(column)
        for row in table.rows:
            for cell, value in zip(row.cells, widths, strict=True):
                tc_width = cell._tc.get_or_add_tcPr().get_or_add_tcW()
                tc_width.set(qn("w:w"), str(value))
                tc_width.set(qn("w:type"), "dxa")

    def _write_pdf(self, path: Path, arguments: CreateDocumentArguments) -> None:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER, TA_LEFT
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import inch
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import (
            ListFlowable,
            ListItem,
            PageBreak,
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
        )

        # A built-in CID font renders Chinese, but some PDF viewers substitute its
        # Latin glyphs with visibly broken spacing. Prefer a real Unicode font and
        # embed it so mixed Chinese/English reports remain portable. The CID font
        # is retained as a last-resort fallback for minimal environments.
        pdf_font_name = "MiniGPT-CJK"
        font_candidates = (
            Path("C:/Windows/Fonts/msyh.ttc"),
            Path("C:/Windows/Fonts/simhei.ttf"),
            Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
            Path("/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc"),
        )
        pdf_font_path = next((item for item in font_candidates if item.is_file()), None)
        if pdf_font_path is not None:
            pdfmetrics.registerFont(TTFont(pdf_font_name, str(pdf_font_path)))
        else:
            pdf_font_name = "STSong-Light"
            pdfmetrics.registerFont(UnicodeCIDFont(pdf_font_name))
        styles = getSampleStyleSheet()
        body = ParagraphStyle(
            "ChineseBody",
            parent=styles["BodyText"],
            fontName=pdf_font_name,
            fontSize=10.5,
            leading=15,
            spaceAfter=8,
            alignment=TA_LEFT,
        )
        heading = ParagraphStyle(
            "ChineseHeading",
            parent=styles["Heading1"],
            fontName=pdf_font_name,
            fontSize=15,
            leading=19,
            textColor=colors.HexColor("#1D4E89"),
            spaceBefore=14,
            spaceAfter=8,
        )
        title_style = ParagraphStyle(
            "ChineseTitle",
            parent=styles["Title"],
            fontName=pdf_font_name,
            fontSize=23,
            leading=29,
            textColor=colors.HexColor("#0F2B4C"),
            alignment=TA_CENTER,
            spaceAfter=8,
        )
        story: list[Any] = [
            Spacer(1, 0.25 * inch),
            Paragraph(arguments.title, title_style),
        ]
        if arguments.subtitle:
            story.append(Paragraph(arguments.subtitle, body))
        story.extend(
            [
                Paragraph(
                    f"作者：{arguments.author}    日期：{date.today().isoformat()}",
                    body,
                ),
                Spacer(1, 0.12 * inch),
                Paragraph("目录", heading),
                ListFlowable(
                    [
                        ListItem(Paragraph(item.heading, body))
                        for item in arguments.sections
                    ],
                    bulletType="1",
                    leftIndent=24,
                ),
            ]
        )
        for item in arguments.sections:
            story.append(Paragraph(item.heading, heading))
            if item.body:
                story.append(Paragraph(item.body.replace("\n", "<br/>"), body))
            if item.bullets:
                story.append(
                    ListFlowable(
                        [ListItem(Paragraph(value, body)) for value in item.bullets],
                        bulletType="bullet",
                        start="-",
                        bulletFontName=pdf_font_name,
                        leftIndent=24,
                    )
                )
        if arguments.metrics:
            story.append(PageBreak())
            story.append(Paragraph("关键指标", heading))
            data = [["指标", "实测值"]] + [
                [row.label, row.value] for row in arguments.metrics
            ]
            table = Table(data, colWidths=[2.5 * inch, 4.0 * inch], repeatRows=1)
            table.setStyle(
                TableStyle(
                    [
                        ("FONTNAME", (0, 0), (-1, -1), pdf_font_name),
                        ("FONTSIZE", (0, 0), (-1, -1), 9.5),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF5")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#0F2B4C")),
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 8),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                        ("TOPPADDING", (0, 0), (-1, -1), 6),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                    ]
                )
            )
            story.append(table)
        if arguments.citations:
            story.append(Paragraph("来源与引用", heading))
            story.append(
                ListFlowable(
                    [ListItem(Paragraph(value, body)) for value in arguments.citations],
                    bulletType="1",
                    leftIndent=24,
                )
            )

        def footer(canvas: Any, document: Any) -> None:
            canvas.saveState()
            canvas.setFont(pdf_font_name, 8)
            canvas.setFillColor(colors.HexColor("#64748B"))
            canvas.drawString(inch, 0.52 * inch, "MiniGPT Agent - 受控文档导出")
            canvas.drawRightString(7.5 * inch, 0.52 * inch, f"第 {document.page} 页")
            canvas.restoreState()

        document = SimpleDocTemplate(
            str(path),
            pagesize=letter,
            rightMargin=inch,
            leftMargin=inch,
            topMargin=inch,
            bottomMargin=0.78 * inch,
            title=arguments.title,
            author=arguments.author,
        )
        document.build(story, onFirstPage=footer, onLaterPages=footer)


def register_document_tools(
    registry: ToolRegistry,
    workspace: DocumentWorkspace,
) -> ToolRegistry:
    """Register six confined document capabilities on an existing registry."""
    registry.register(
        ToolSpec(
            name="create_docx",
            description="Create a styled Word .docx inside the configured workspace.",
            arguments_model=CreateDocumentArguments,
            handler=workspace.create_docx,
            timeout_seconds=20,
            side_effect=True,
        )
    )
    registry.register(
        ToolSpec(
            name="create_pdf",
            description="Create a styled PDF inside the configured workspace.",
            arguments_model=CreateDocumentArguments,
            handler=workspace.create_pdf,
            timeout_seconds=20,
            side_effect=True,
        )
    )
    registry.register(
        ToolSpec(
            name="read_docx",
            description="Extract paragraphs and table text from a workspace DOCX.",
            arguments_model=ReadDocumentArguments,
            handler=workspace.read_docx,
        )
    )
    registry.register(
        ToolSpec(
            name="read_pdf",
            description="Extract text from all pages of a workspace PDF.",
            arguments_model=ReadDocumentArguments,
            handler=workspace.read_pdf,
        )
    )
    registry.register(
        ToolSpec(
            name="summarize_document",
            description="Create a bounded extractive summary of a workspace DOCX or PDF.",
            arguments_model=SummarizeDocumentArguments,
            handler=workspace.summarize_document,
        )
    )
    registry.register(
        ToolSpec(
            name="export_report",
            description="Export structured metrics and findings as a formal DOCX or PDF.",
            arguments_model=ExportReportArguments,
            handler=workspace.export_report,
            timeout_seconds=20,
            side_effect=True,
        )
    )
    return registry
